"""Testes do dashboard Flask - rotas e lógica de negócio."""

import pytest
import sys
import os
import json


class TestLogin:
    """Testes de autenticação."""

    def test_login_page_renders(self, client):
        response = client.get('/login')
        assert response.status_code == 200

    def test_login_with_invalid_credentials(self, client, app):
        with app.app_context():
            from app.models import db, User
            db.create_all()
        
        response = client.post('/login', data={
            'username': 'admin',
            'password': 'wrong_password'
        }, follow_redirects=True)
        assert response.status_code == 200
        
        with client.session_transaction() as sess:
            assert 'user' not in sess

    def test_protected_route_redirects_without_session(self, client):
        response = client.get('/', follow_redirects=False)
        assert response.status_code == 302
        assert '/login' in response.headers['Location']

    def test_index_page_renders_with_session(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        response = auth_client.get('/')
        assert response.status_code == 200
        assert b"Registrale" in response.data or b"html" in response.data.lower()


class TestCompanies:
    """Testes de empresas."""

    def test_get_companies_empty(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        response = auth_client.get('/api/companies')
        assert response.status_code == 200
        data = response.get_json()
        assert isinstance(data, list)

    def test_create_company(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        
        response = auth_client.post('/api/companies', json={
            'nome': 'Empresa Teste',
            'cnpj': '12.345.678/0001-90',
            'uf': 'SP',
            'cidade': 'São Paulo',
            'email': 'teste@teste.com',
            'telefone': '11999999999',
            'situacao': 'Ativa',
            'status': True,
            'origem': 'Manual'
        })
        assert response.status_code == 200
        data = response.get_json()
        assert data.get('status') == 'success' or data.get('id') is not None

    def test_create_company_duplicate_cnpj(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company
            db.create_all()
            existing = Company(
                nome="Existente",
                cnpj="99.999.999/0001-99",
                cnpj_norm="99999999000199"
            )
            db.session.add(existing)
            db.session.commit()
        
        response = auth_client.post('/api/companies', json={
            'nome': 'Nova Empresa',
            'cnpj': '99.999.999/0001-99'
        })
        assert response.status_code in [400, 409]

    def test_create_company_missing_fields(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        
        response = auth_client.post('/api/companies', json={
            'nome': 'Empresa Sem CNPJ'
        })
        assert response.status_code >= 400

    def test_create_company_invalid_cnpj(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        
        response = auth_client.post('/api/companies', json={
            'nome': 'Empresa CNPJ Inválido',
            'cnpj': 'CNPJ_INVALIDO'
        })
        assert response.status_code >= 400

    def test_unmonitor_by_origin_endpoint(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company
            db.create_all()
            db.session.add(Company(nome="Empresa GC", cnpj="11.111.111/0001-11", cnpj_norm="11111111000111", origem="GestãoClick", status=True))
            db.session.add(Company(nome="Empresa Planilha", cnpj="22.222.222/0001-22", cnpj_norm="22222222000122", origem="Planilha", status=True))
            db.session.add(Company(nome="Empresa Manual", cnpj="33.333.333/0001-33", cnpj_norm="33333333000133", origem="Manual", status=True))
            db.session.commit()
            
        resp = auth_client.post('/api/companies/unmonitor_by_origin', json={
            "origins": ["GestãoClick", "Planilha"]
        })
        assert resp.status_code == 200
    def test_toggle_origin_monitoring_endpoint(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company
            db.create_all()
            db.session.add(Company(nome="Empresa GC Test", cnpj="44.444.444/0001-44", cnpj_norm="44444444000144", origem="GestãoClick", status=True))
            db.session.commit()
            
        # Desativar monitoramento da origem GestãoClick
        resp = auth_client.post('/api/companies/toggle_origin_monitoring', json={
            "origin": "GestãoClick",
            "status": False
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "success"
        
        with app.app_context():
            from app.models import Company
            c = Company.query.filter_by(cnpj_norm="44444444000144").first()
            assert c.status is False
            
        # Reativar monitoramento da origem GestãoClick
        resp2 = auth_client.post('/api/companies/toggle_origin_monitoring', json={
            "origin": "GestãoClick",
            "status": True
        })
        assert resp2.status_code == 200
        with app.app_context():
            from app.models import Company
            c = Company.query.filter_by(cnpj_norm="44444444000144").first()
            assert c.status is True


class TestTemplates:
    """Testes de templates de email."""

    def test_get_templates_empty(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        response = auth_client.get('/api/templates')
        assert response.status_code == 200
        data = response.get_json()
        assert isinstance(data, list)

    def test_create_template(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        
        response = auth_client.post('/api/templates', json={
            'name': 'Template Teste',
            'subject': 'Assunto Teste',
            'body_html': '<h1>Teste</h1>'
        })
        assert response.status_code == 200

    def test_cannot_delete_default_template(self, auth_client, app):
        with app.app_context():
            from app.models import db, EmailTemplate
            template = EmailTemplate.query.filter_by(name='Padrão Registrale').first()
            template_id = template.id
        
        response = auth_client.delete(f'/api/templates/{template_id}')
        assert response.status_code == 400

    def test_create_template_missing_data(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
        
        response = auth_client.post('/api/templates', json={
            'name': 'Sem Body HTML'
        })
        assert response.status_code >= 400

    def test_delete_nonexistent_template(self, auth_client, app):
        with app.app_context():
            from app.models import db
            db.create_all()
            
        response = auth_client.delete('/api/templates/999')
        assert response.status_code >= 400

    def test_apply_highlight_to_trecho_and_email_rendering(self, app):
        with app.app_context():
            from app.services.email_service import apply_highlight_to_trecho, build_mentions_email_html
            
            # 1. Teste de destaque de CNPJ e termos
            raw_text = "Contrato celebrado com 12.345.678/0001-90 para prestação de serviços."
            highlighted = apply_highlight_to_trecho(raw_text, cnpj="12.345.678/0001-90")
            assert "highlight" in highlighted
            assert "#FFA" in highlighted
            assert "12.345.678/0001-90" in highlighted
            
            # 2. Teste de placeholders <%%>...</%%>
            placeholder_text = "Aviso de <%%>LICITAÇÃO</%%> referente ao pregão."
            h2 = apply_highlight_to_trecho(placeholder_text)
            assert "LICITAÇÃO" in h2
            assert "background-color: #FFA" in h2
            assert "<%%>" not in h2
            
            # 3. Teste de build_mentions_email_html
            mentions = [{
                'id': 'm1',
                'empresa': 'EMPRESA TESTE',
                'cnpj': '12.345.678/0001-90',
                'secao': 'SECAO_3',
                'data': '2026-08-23',
                'trecho': 'Publicação da EMPRESA TESTE 12.345.678/0001-90 no DOU.',
                'link': 'https://in.gov.br/teste'
            }]
            html = build_mentions_email_html(mentions)
            assert "EMPRESA TESTE" in html
            assert "background-color: #FFA" in html
            assert "SECAO_3" in html


class TestClearData:
    """Testes de limpeza de dados."""

    def test_clear_mentions(self, auth_client, app):
        with app.app_context():
            from app.models import db, Mention
            db.create_all()
            mention = Mention(id="test1", empresa="Teste", cnpj="123")
            db.session.add(mention)
            db.session.commit()
        
        response = auth_client.post('/api/admin/clear_data', json={'type': 'mentions'})
        assert response.status_code == 200
        
        with app.app_context():
            from app.models import Mention as M
            assert M.query.count() == 0

    def test_clear_requires_master_role(self, user_client, app):
        response = user_client.post('/api/admin/clear_data', json={'type': 'mentions'})
        assert response.status_code == 403


class TestNewOptimizedFeatures:
    """Testes para as novas funcionalidades de busca, toggle de rotinas e exportações com abas."""

    def test_search_companies_endpoint(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company
            db.create_all()
            c = Company(nome="Acme Corporacao", cnpj="77.888.999/0001-11", cnpj_norm="77888999000111", origem="Manual", status=True)
            db.session.add(c)
            db.session.commit()

        # Busca por nome
        resp = auth_client.get('/api/companies/search?q=Acme')
        assert resp.status_code == 200
        data = resp.get_json()
        assert len(data) >= 1
        assert "Acme" in data[0]["nome"]

        # Busca por CNPJ
        resp_cnpj = auth_client.get('/api/companies/search?q=77888999')
        assert resp_cnpj.status_code == 200
        data_cnpj = resp_cnpj.get_json()
        assert len(data_cnpj) >= 1

    def test_sheets_cnpjs_endpoint(self, auth_client, app):
        with app.app_context():
            from app.models import db, Settings
            db.create_all()
            s = Settings.query.filter_by(key='global_settings').first()
            if not s:
                s = Settings(key='global_settings')
                db.session.add(s)
            s.set_value({"google_sheets": {}})
            db.session.commit()

        resp = auth_client.get('/api/sheets/cnpjs')
        assert resp.status_code == 200
        data = resp.get_json()
        assert "status" in data

    def test_export_report_excel_has_metadata_sheet(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company
            db.create_all()
            db.session.add(Company(nome="Empresa Export", cnpj="11.222.333/0001-44", cnpj_norm="11222333000144", origem="Manual", status=True))
            db.session.commit()

        resp = auth_client.get('/api/export_report')
        assert resp.status_code == 200
        assert resp.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        # Carrega o workbook retornado para verificar as 2 abas
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        sheet_names = wb.sheetnames
        assert "Empresas Monitoradas" in sheet_names
        assert "Metadados da Geração" in sheet_names

    def test_export_pdf_endpoint(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company
            db.create_all()
            db.session.add(Company(nome="Empresa PDF", cnpj="22.333.444/0001-55", cnpj_norm="22333444000155", origem="Manual", status=True))
            db.session.commit()

        resp = auth_client.post('/api/export_pdf', json={
            'companies': [{'nome': 'Empresa PDF', 'cnpj': '22.333.444/0001-55', 'origem': 'Manual', 'status': True}]
        })
        assert resp.status_code == 200
        assert resp.headers["Content-Type"] == "application/pdf"
        assert len(resp.data) > 1000

    def test_download_missing_inlabs_endpoint(self, auth_client):
        resp = auth_client.post('/api/routines/download_missing_inlabs', json={
            'days': ['2026-08-01', '2026-08-02'],
            'month': 8,
            'year': 2026
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "success"
        assert "2 dia(s)" in data["message"]

    def test_check_date_endpoint(self, auth_client):
        resp = auth_client.get('/api/routines/check_date?date=2026-08-15')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "success"
        assert data["date"] == "2026-08-15"
        assert "already_loaded" in data

    def test_trigger_paused_routine_blocked(self, auth_client, tmp_path, monkeypatch):
        import yaml
        from app.services import dag_config_service
        
        # Cria rotina pausada temporária
        paused_yaml = {
            "dag": {
                "id": "teste_pausado",
                "active": False,
                "search": [{"terms": ["TESTE"]}]
            }
        }
        routine_file = tmp_path / "teste_pausado.yaml"
        with open(routine_file, "w", encoding="utf-8") as f:
            yaml.safe_dump(paused_yaml, f)
            
        monkeypatch.setattr(dag_config_service, "get_dag_confs_path", lambda: str(tmp_path))
        from app.routes import dags
        monkeypatch.setattr(dags, "get_dag_confs_path", lambda: str(tmp_path))
        
        resp = auth_client.post('/api/routines/trigger/teste_pausado.yaml', json={})
        assert resp.status_code == 400
        data = resp.get_json()
        assert "pausada" in data["message"].lower()

    def test_trigger_monthly_paused_routine_blocked(self, auth_client, tmp_path, monkeypatch):
        import yaml
        from app.services import dag_config_service
        
        paused_yaml = {
            "dag": {
                "id": "teste_mensal_pausado",
                "active": False,
                "search": [{"terms": ["TESTE"]}]
            }
        }
        routine_file = tmp_path / "teste_mensal_pausado.yaml"
        with open(routine_file, "w", encoding="utf-8") as f:
            yaml.safe_dump(paused_yaml, f)
            
        monkeypatch.setattr(dag_config_service, "get_dag_confs_path", lambda: str(tmp_path))
        from app.routes import dags
        monkeypatch.setattr(dags, "get_dag_confs_path", lambda: str(tmp_path))
        
        resp = auth_client.post('/api/routines/trigger_monthly', json={
            "year": 2026,
            "month": 8,
            "routines": ["teste_mensal_pausado.yaml"]
        })
        assert resp.status_code == 400
        data = resp.get_json()
        assert "pausada" in data["message"].lower()

    def test_toggle_routine_endpoint(self, auth_client, tmp_path, monkeypatch):
        import yaml
        from app.services import dag_config_service
        from app.routes import dags
        from app.services import airflow_service
        
        test_yaml = {
            "dag": {
                "id": "teste_toggle",
                "active": True,
                "search": [{"terms": ["TESTE"]}]
            }
        }
        routine_file = tmp_path / "teste_toggle.yaml"
        with open(routine_file, "w", encoding="utf-8") as f:
            yaml.safe_dump(test_yaml, f)
            
        monkeypatch.setattr(dag_config_service, "get_dag_confs_path", lambda: str(tmp_path))
        monkeypatch.setattr(dags, "get_dag_confs_path", lambda: str(tmp_path))
        monkeypatch.setattr(airflow_service, "toggle_airflow_dag", lambda *args, **kwargs: True)
        
        # Pausar rotina
        resp = auth_client.post('/api/routines/toggle/teste_toggle.yaml', json={"active": False})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "success"
        assert data["active"] is False
        
        with open(routine_file, "r", encoding="utf-8") as f:
            saved = yaml.safe_load(f)
            assert saved["dag"]["active"] is False
            assert saved["dag"]["is_paused"] is True
            
        # Reativar rotina
        resp = auth_client.post('/api/routines/toggle/teste_toggle.yaml', json={"active": True})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "success"
        assert data["active"] is True
        
        with open(routine_file, "r", encoding="utf-8") as f:
            saved = yaml.safe_load(f)
            assert saved["dag"]["active"] is True
            assert saved["dag"]["is_paused"] is False

    def test_companies_endpoint_returns_total_mencoes(self, auth_client, app):
        with app.app_context():
            from app.models import db, Company, Mention
            db.create_all()
            db.session.add(Company(nome="Empresa Com Mencoes", cnpj="55.666.777/0001-88", cnpj_norm="55666777000188", origem="Manual", status=True))
            db.session.add(Mention(id="men_test_1", empresa="Empresa Com Mencoes", cnpj="55.666.777/0001-88", cnpj_norm="55666777000188", data="2026-08-20", secao="1", trecho="Trecho teste"))
            db.session.add(Mention(id="men_test_2", empresa="Empresa Com Mencoes", cnpj="55.666.777/0001-88", cnpj_norm="55666777000188", data="2026-08-21", secao="1", trecho="Trecho teste 2"))
            db.session.commit()

        resp = auth_client.get('/api/companies')
        assert resp.status_code == 200
        companies = resp.get_json()
        target = next((c for c in companies if c['cnpj_norm'] == '55666777000188'), None)
        assert target is not None
        assert target.get('total_mencoes') == 2

    def test_inlabs_retention_limit_purges_oldest_and_protects_target(self, app, monkeypatch):
        with app.app_context():
            from app.models import db, InlabsDownloadLog
            from app.services import inlabs_service
            db.create_all()
            
            # Limpa logs anteriores
            InlabsDownloadLog.query.delete()
            
            # Cria 5 dias baixados com timestamps diferentes
            dates = [
                ('2026-04-01', '2026-04-01 10:00:00'), # Mais antigo
                ('2026-04-02', '2026-04-02 10:00:00'),
                ('2026-04-03', '2026-04-03 10:00:00'),
                ('2000-01-01', '2026-08-23 11:00:00'), # Data histórica (ano 2000) mas baixada recentemente
                ('2000-01-02', '2026-08-23 11:05:00'),
            ]
            for dt, dl_at in dates:
                db.session.add(InlabsDownloadLog(date_str=dt, downloaded_at=dl_at, status='success'))
            db.session.commit()
            
            # Simula que o PostgreSQL possui exatamente essas 5 datas
            monkeypatch.setattr(inlabs_service, "get_downloaded_dates", lambda *args, **kwargs: set(d[0] for d in dates))
            monkeypatch.setattr(inlabs_service, "get_inlabs_postgres_engine", lambda: None)
            
            # Executa com limite máximo de 3 dias, protegendo '2000-01-01' e '2000-01-02'
            # Deve apagar os 2 dias baixados mais antigamente ('2026-04-01' e '2026-04-02')
            deleted = inlabs_service.enforce_inlabs_retention_limit(max_days=3, protected_dates=['2000-01-01', '2000-01-02'])
            
            # Verifica que no SQLite restaram 3 datas e as datas do ano 2000 continuam intactas
            remaining = [l.date_str for l in InlabsDownloadLog.query.all()]
            assert '2000-01-01' in remaining
            assert '2000-01-02' in remaining
            assert '2026-04-01' not in remaining

    def test_record_inlabs_download_success(self, app):
        with app.app_context():
            from app.models import db, InlabsDownloadLog
            from app.services.inlabs_service import record_inlabs_download_success
            db.create_all()
            
            record_inlabs_download_success('2026-08-23')
            log = InlabsDownloadLog.query.filter_by(date_str='2026-08-23').first()
            assert log is not None
            assert log.status == 'success'
            assert log.downloaded_at is not None

    def test_wait_for_specific_dag_runs(self, monkeypatch):
        import requests
        from app.services.airflow_service import wait_for_specific_dag_runs
        
        class MockResponse:
            def __init__(self, state):
                self.status_code = 200
                self._state = state
            def json(self):
                return {"state": self._state}
                
        def mock_get(url, **kwargs):
            return MockResponse("success")
            
        monkeypatch.setattr(requests, "get", mock_get)
        all_finished, states = wait_for_specific_dag_runs("ro-dou_inlabs_load_pg", ["run_1", "run_2"], max_wait=5, poll_interval=1)
        assert all_finished is True
        assert states.get("run_1") == "success"
        assert states.get("run_2") == "success"

    def test_holiday_service_excludes_holidays_and_weekends(self):
        from app.services.holiday_service import get_business_days_for_month, is_business_day, is_within_inlabs_retention_window
        from datetime import date
        
        # 1 de janeiro de 2026 é quinta-feira (Ano Novo) - deve ser excluído
        is_bus, reason = is_business_day(date(2026, 1, 1))
        assert is_bus is False
        assert "Ano Novo" in reason or "Confraternização" in reason

        # 2 de janeiro de 2026 é sexta-feira - dia útil válido
        is_bus, reason = is_business_day(date(2026, 1, 2))
        assert is_bus is True

        # Janeiro de 2026 completo
        biz_days = get_business_days_for_month(2026, 1, cap_today=False)
        assert '2026-01-01' not in biz_days
        assert '2026-01-02' == biz_days[0]
        # Finais de semana não devem estar presentes
        assert '2026-01-03' not in biz_days # Sábado
        assert '2026-01-04' not in biz_days # Domingo

        # Feriado móvel: Sexta-feira Santa em 2026 (03/04/2026)
        is_bus_santa, _ = is_business_day(date(2026, 4, 3))
        assert is_bus_santa is False

        # Tiradentes (21/04/2026)
        is_bus_tiradentes, _ = is_business_day(date(2026, 4, 21))
        assert is_bus_tiradentes is False

        # Consciência Negra (20/11/2026)
        is_bus_zumbi, _ = is_business_day(date(2026, 11, 20))
        assert is_bus_zumbi is False

        # Janela de 120 dias do INLABS
        # Data de 1 ano atrás deve estar fora da janela
        assert is_within_inlabs_retention_window('2020-01-01', 120) is False

    def test_monthly_inlabs_check_endpoint_with_holidays_and_120_days_rule(self, auth_client):
        resp = auth_client.get('/api/routines/monthly_inlabs_check?month=1&year=2026&routine=Pesquisa_cnpj.yaml')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'ok'
        assert '2026-01-01' not in data.get('missing_days')
        assert '2026-01-01' not in data.get('inlabs_days')
        assert 'api_dou_days' in data
        assert 'downloadable_inlabs_days' in data
        assert 'scenario' in data
        assert data.get('scenario') in ['all_api_dou', 'mixed', 'download_only', 'complete']

    def test_trigger_monthly_executes_api_dou_for_historical_missing_days(self, auth_client, app, monkeypatch):
        from app.routes import dags
        
        triggered_calls = []
        def mock_trigger(dag_id, logical_date=None, **kwargs):
            triggered_calls.append((dag_id, logical_date))
            return True, "Triggered", {"dag_id": dag_id, "dag_run_id": f"run_{dag_id}_{logical_date}"}
            
        monkeypatch.setattr(dags, "trigger_airflow_dag", mock_trigger)
        monkeypatch.setattr(dags, "fetch_mentions_from_dag_run", lambda *args, **kwargs: [])
        
        resp = auth_client.post('/api/routines/trigger_monthly', json={
            "month": 1,
            "year": 2026,
            "routines": ["Pesquisa_cnpj.yaml"],
            "mode": "download_and_search"
        })
        assert resp.status_code == 200
        assert resp.get_json().get("status") == "success"

    def test_trigger_monthly_api_dou_only_mode(self, auth_client, monkeypatch):
        from app.routes import dags
        
        triggered_calls = []
        def mock_trigger(dag_id, logical_date=None, **kwargs):
            triggered_calls.append((dag_id, logical_date))
            return True, "Triggered", {"dag_id": dag_id, "dag_run_id": f"run_{dag_id}_{logical_date}"}
            
        monkeypatch.setattr(dags, "trigger_airflow_dag", mock_trigger)
        monkeypatch.setattr(dags, "fetch_mentions_from_dag_run", lambda *args, **kwargs: [])
        
        resp = auth_client.post('/api/routines/trigger_monthly', json={
            "month": 4,
            "year": 2026,
            "routines": ["Pesquisa_cnpj.yaml"],
            "mode": "api_dou_only"
        })
        assert resp.status_code == 200
        assert resp.get_json().get("status") == "success"

    def test_monthly_inlabs_check_scenario_mixed_with_downloadable_zero(self, auth_client, monkeypatch):
        from app.services import inlabs_service
        # Simula que a partir do dia 24/04/2026 está baixado no banco INLABS
        monkeypatch.setattr(inlabs_service, "get_downloaded_dates", lambda f, l: ['2026-04-24', '2026-04-27', '2026-04-28', '2026-04-29', '2026-04-30'])
        
        resp = auth_client.get('/api/routines/monthly_inlabs_check?month=4&year=2026&routine=Pesquisa_cnpj.yaml')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'ok'
        assert data.get('inlabs_count') > 0
        assert data.get('api_dou_count') > 0
        assert data.get('scenario') == 'mixed'

    def test_check_date_returns_holiday_and_120_info(self, auth_client):
        # 01/01/2026 é Confraternização Universal (feriado) e fora de 120 dias
        resp = auth_client.get('/api/routines/check_date?date=2026-01-01')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        assert data.get('is_business_day') is False
        assert 'Confraternização Universal' in data.get('holiday_reason')
        assert data.get('is_within_120') is False

    def test_trigger_individual_routine_historical_date_executes_api_dou(self, auth_client, monkeypatch):
        from app.routes import dags
        
        triggered_calls = []
        def mock_trigger(dag_id, logical_date=None, **kwargs):
            triggered_calls.append((dag_id, logical_date))
            return True, "Triggered", {"dag_id": dag_id, "dag_run_id": f"run_{dag_id}_{logical_date}"}
            
        monkeypatch.setattr(dags, "trigger_airflow_dag", mock_trigger)
        monkeypatch.setattr(dags, "fetch_mentions_from_dag_run", lambda *args, **kwargs: [])
        
        resp = auth_client.post('/api/routines/trigger/Pesquisa_cnpj.yaml', json={
            "logical_date": "2026-01-15"
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        assert any(k in data.get('message', '').lower() for k in ["disparada", "sucesso", "api oficial"])

    def test_trigger_individual_routine_with_brazilian_date_format(self, auth_client, monkeypatch):
        from app.routes import dags
        from app.services import inlabs_service
        
        triggered_calls = []
        def mock_trigger(dag_id, logical_date=None, **kwargs):
            triggered_calls.append((dag_id, logical_date))
            return True, "Triggered", {"dag_id": dag_id, "dag_run_id": f"run_{dag_id}_{logical_date}"}
            
        monkeypatch.setattr(dags, "trigger_airflow_dag", mock_trigger)
        monkeypatch.setattr(dags, "fetch_mentions_from_dag_run", lambda *args, **kwargs: [])
        monkeypatch.setattr(inlabs_service, "is_date_loaded", lambda d: (True, 50))
        
        # Envia a data em formato brasileiro DD/MM/AAAA
        resp = auth_client.post('/api/routines/trigger/Pesquisa_cnpj.yaml', json={
            "logical_date": "15/01/2026"
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        # Verifica que a data foi normalizada para YYYY-MM-DD
        assert len(triggered_calls) > 0
        assert triggered_calls[0][1] == "2026-01-15"

    def test_cleanup_orphaned_temp_dags(self, app):
        import os
        from app.services.dag_config_service import get_dag_confs_path, cleanup_orphaned_temp_dags
        dag_confs_path = get_dag_confs_path()
        
        # Cria um arquivo temporário de teste
        temp_file = os.path.join(dag_confs_path, "temp_adhoc_dou_unittest.yaml")
        with open(temp_file, 'w', encoding='utf-8') as f:
            f.write("dag:\n  id: temp_adhoc_dou_unittest\n  schedule: null\n")
            
        assert os.path.exists(temp_file)
        cleaned = cleanup_orphaned_temp_dags(max_age_seconds=0, force_all=True)
        assert cleaned >= 1
        assert not os.path.exists(temp_file)

    def test_cleanup_temp_dags_endpoint(self, auth_client):
        resp = auth_client.post('/api/routines/cleanup_temp')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        assert 'cleaned_count' in data

    def test_get_user_manual_endpoint(self, auth_client):
        resp = auth_client.get('/api/manual')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        assert 'Manual do Usuário' in data.get('content')
        assert 'Visão Geral' in data.get('content')

    def test_download_user_manual_endpoint(self, auth_client):
        resp = auth_client.get('/api/manual?download=true')
        assert resp.status_code == 200
        assert b'# Manual do Usu' in resp.data
        assert 'text/markdown' in resp.content_type

    def test_main_dag_status_endpoint(self, auth_client):
        resp = auth_client.get('/api/system/main_dag_status')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'ok'
        assert 'is_configured' in data
        assert 'missing_fields' in data
        assert 'main_dag' in data
        assert 'smtp' in data
        assert 'smtp_configured' in data

    def test_configure_main_dag_endpoint(self, auth_client, app):
        resp = auth_client.post('/api/system/configure_main_dag', json={
            "emails": "alerta@empresa.com, diretoria@empresa.com",
            "subject": "[Monitoramento DOU] Alerta Diário",
            "schedule": "0 8 * * MON-FRI",
            "smtp": {
                "server": "smtp.mailgun.org",
                "port": 587,
                "user": "postmaster@empresa.com",
                "password": "secretpassword",
                "from_email": "dou@empresa.com"
            }
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        
        # Verify status endpoint returns configured
        status_resp = auth_client.get('/api/system/main_dag_status')
        assert status_resp.status_code == 200
        status_data = status_resp.get_json()
        assert status_data.get('is_configured') is True
        assert "alerta@empresa.com" in status_data.get('main_dag', {}).get('emails', [])
        assert status_data.get('main_dag', {}).get('subject') == "[Monitoramento DOU] Alerta Diário"

    def test_create_routine_defaults_to_inlabs(self, auth_client, app):
        import os
        from app.services.dag_config_service import get_dag_confs_path
        
        routine_name = "Teste Rotina INLABS Padrao"
        resp = auth_client.post('/api/routines', json={
            "id": "teste_inlabs_padrao",
            "file": "teste_inlabs_padrao.yaml",
            "name": routine_name,
            "terms": ["termo_teste_inlabs"],
            "sections": ["SECAO_1", "SECAO_2", "SECAO_3"],
            "emails": ["teste@inlabs.com"]
            # Sem passar source explicitamente, deve assumir INLABS
        })
        assert resp.status_code == 200
        
        # Carrega a lista de rotinas e verifica a fonte
        routines_resp = auth_client.get('/api/routines')
        assert routines_resp.status_code == 200
        routines = routines_resp.get_json()
        created = next((r for r in routines if r.get('id') == 'teste_inlabs_padrao' or r.get('file') == 'teste_inlabs_padrao.yaml'), None)
        assert created is not None
        assert created.get('source') == 'INLABS'
        
        # Limpa arquivo gerado
        yaml_file = os.path.join(get_dag_confs_path(), "teste_inlabs_padrao.yaml")
        if os.path.exists(yaml_file):
            try:
                os.remove(yaml_file)
            except Exception:
                pass

    def test_get_settings_endpoint(self, auth_client):
        resp = auth_client.get('/api/settings')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'ok'
        assert 'settings' in data
        assert 'smtp' in data['settings']

    def test_save_settings_preserves_password(self, auth_client, app):
        # 1. Salva com senha
        resp1 = auth_client.post('/api/save_settings', json={
            "smtp": {
                "server": "smtp.gmail.com",
                "port": "587",
                "user": "teste@gmail.com",
                "password": "MinhaSenhaSuperSecreta123",
                "from_email": "teste@gmail.com"
            }
        })
        assert resp1.status_code == 200
        
        # 2. Salva novamente sem enviar senha (string vazia)
        resp2 = auth_client.post('/api/save_settings', json={
            "smtp": {
                "server": "smtp.gmail.com",
                "port": "587",
                "user": "teste@gmail.com",
                "password": "",
                "from_email": "teste@gmail.com"
            }
        })
        assert resp2.status_code == 200
        
        # 3. Verifica se a senha foi preservada
        from app.models import Settings
        with app.app_context():
            s = Settings.query.filter_by(key='global_settings').first()
            val = s.get_value()
            assert val['smtp']['password'] == "MinhaSenhaSuperSecreta123"

    def test_google_sheets_diagnosis_with_spreadsheet_url(self, auth_client):
        # Configura Google Sheets com spreadsheet_url e credentials_json
        auth_client.post('/api/save_settings', json={
            "google_sheets": {
                "spreadsheet_url": "https://docs.google.com/spreadsheets/d/123456789/edit",
                "credentials_json": '{"type": "service_account", "client_email": "test@service.iam.gserviceaccount.com"}'
            }
        })
        
        resp = auth_client.get('/api/system/integrations_status')
        assert resp.status_code == 200
        data = resp.get_json()
        gs_item = next((i for i in data.get('integrations', []) if i['id'] == 'google_sheets'), None)
        assert gs_item is not None
        assert gs_item['is_configured'] is True

    def test_test_smtp_fallback_to_saved_password(self, auth_client):
        # Configura SMTP com senha no banco
        auth_client.post('/api/save_settings', json={
            "smtp": {
                "server": "localhost",
                "port": "2525",
                "user": "test_user",
                "password": "saved_password",
                "from_email": "test_user@localhost"
            }
        })
        
        # Testa chamada com senha em branco no payload
        # (vai falhar na conexão de rede com localhost:2525, mas não por falta de parâmetros)
        resp = auth_client.post('/api/test_smtp', json={
            "smtp": {
                "server": "localhost",
                "port": "2525",
                "user": "test_user",
                "password": ""
            },
            "test_email": "destino@teste.com"
        })
        # Não deve falhar por falta de parâmetros (como senha/servidor), e sim tentar a conexão de rede
        data = resp.get_json()
        assert data.get('message') != "Servidor SMTP, porta e email de teste são obrigatórios."
        assert "conexão" in data.get('message', '').lower() or "sucesso" in data.get('message', '').lower()

    def test_sync_gestaoclick_without_credentials_returns_400(self, auth_client, app, monkeypatch):
        # Garante que não há variáveis de ambiente nem chaves no banco
        monkeypatch.delenv("ACCESS_TOKEN", raising=False)
        monkeypatch.delenv("SECRET_ACCESS_TOKEN", raising=False)
        with app.app_context():
            from app.models import db, Settings
            s = Settings.query.filter_by(key='global_settings').first()
            if s:
                val = s.get_value()
                val['api_keys'] = {}
                s.set_value(val)
                db.session.commit()

        resp = auth_client.post('/api/sync')
        assert resp.status_code == 400
        data = resp.get_json()
        assert data.get('status') == 'error'
        assert 'Credenciais do GestãoClick não configuradas' in data.get('message', '')

    def test_sync_gestaoclick_with_credentials_starts(self, auth_client, app):
        with app.app_context():
            from app.models import db, Settings
            s = Settings.query.filter_by(key='global_settings').first()
            if not s:
                s = Settings(key='global_settings')
                db.session.add(s)
            val = s.get_value() if s.value else {}
            val['api_keys'] = {
                'gestaoclick_access_token': 'fake_token_123',
                'gestaoclick_secret_token': 'fake_secret_456',
                'gestaoclick_base_url': 'https://api.gestaoclick.com/franquias'
            }
            s.set_value(val)
            db.session.commit()

        resp = auth_client.post('/api/sync')
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'

    def test_export_mentions_pdf_and_excel_with_section_filter(self, auth_client):
        sample_mentions = [
            {
                "id": "m1",
                "empresa": "Empresa Teste 1",
                "cnpj": "11.222.333/0001-44",
                "secao": "DOU - Seção 1",
                "data": "20/08/2026",
                "trecho": "Trecho de teste com publicação",
                "link": "https://in.gov.br/m1"
            },
            {
                "id": "m2",
                "empresa": "Empresa Teste 2",
                "cnpj": "22.333.444/0001-55",
                "secao": "DOU - Seção 3",
                "data": "21/08/2026",
                "trecho": "Trecho de teste seção 3",
                "link": "https://in.gov.br/m2"
            }
        ]
        
        # Test PDF export with section filter
        pdf_resp = auth_client.post('/api/export_mentions_pdf', json={
            "mentions": sample_mentions,
            "filters": {"section": "SECAO_1"}
        })
        assert pdf_resp.status_code == 200
        assert pdf_resp.mimetype == 'application/pdf'
        
        # Test Excel export with section filter
        excel_resp = auth_client.post('/api/export_mentions_excel', json={
            "mentions": sample_mentions,
            "filters": {"section": "SECAO_1"}
        })
        assert excel_resp.status_code == 200
        assert 'spreadsheet' in excel_resp.mimetype or 'openxmlformats' in excel_resp.mimetype or len(excel_resp.data) > 0

    def test_save_main_routine_exact_search_false_persists(self, auth_client, app):
        import yaml
        from app.models import db, Settings, Company
        from app.services.dag_config_service import get_dag_confs_path, rebuild_yaml_from_db, get_base_yaml_path

        # 1. Salva a rotina padrão com busca por termo exato DESATIVADA (is_exact_search=False)
        resp = auth_client.post('/api/routines', json={
            "file": "Pesquisa_cnpj.yaml",
            "name": "MONITORAMENTO PADRAO TESTE",
            "schedule": "0 8 * * MON-FRI",
            "terms": ["12345678000190"],
            "organs": ["ANVISA"],
            "sections": ["SECAO_1", "SECAO_2"],
            "emails": ["notificacao@empresa.com"],
            "subject": "[Monitoramento] Alertas DOU",
            "active": True,
            "is_exact_search": False,
            "source": "INLABS"
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'

        # 2. Verifica se persistiu no SQLite
        with app.app_context():
            s_rec = Settings.query.filter_by(key='main_dag_settings').first()
            assert s_rec is not None
            val = s_rec.get_value()
            assert val.get('is_exact_search') is False

        # 3. Verifica se a API /api/routines retorna is_exact_search=False
        routines_resp = auth_client.get('/api/routines')
        assert routines_resp.status_code == 200
        routines = routines_resp.get_json()
        main_r = next((r for r in routines if r.get('file') == 'Pesquisa_cnpj.yaml' or r.get('type') == 'sync'), None)
        assert main_r is not None
        assert main_r.get('is_exact_search') is False

        # 4. Força uma reconstrução (simulando cadastro de nova empresa ou sincronização)
        with app.app_context():
            new_c = Company(nome="Empresa Rebuild Test", cnpj="33.444.555/0001-66", cnpj_norm="33444555000166", status=True)
            db.session.add(new_c)
            db.session.commit()
            rebuild_yaml_from_db()

        # 5. Verifica se no arquivo YAML gerado a busca exata continua False
        base_yaml = get_base_yaml_path()
        assert os.path.exists(base_yaml)
        with open(base_yaml, 'r', encoding='utf-8') as f:
            ydata = yaml.safe_load(f)
            searches = ydata.get('dag', {}).get('search', [])
            assert len(searches) > 0
            assert searches[0].get('is_exact_search') is False

    def test_send_email_success_response(self, auth_client, monkeypatch):
        from unittest.mock import MagicMock
        from app.services import email_service

        # Mock da chamada SMTP dentro de EmailSender
        mock_sender = MagicMock()
        mock_sender.send_custom_email.return_value = True
        monkeypatch.setattr(email_service, "EmailSender", lambda *args, **kwargs: mock_sender)

        resp = auth_client.post('/api/send_email', json={
            "to_emails": ["destinatario@teste.com", "outro@teste.com"],
            "subject": "Relatório de Teste",
            "body_html": "<p>Conteúdo de Teste</p>"
        })
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('status') == 'success'
        assert 'sucesso' in data.get('message', '').lower()

    def test_email_sender_service_returns_true(self, monkeypatch):
        import smtplib
        from unittest.mock import MagicMock
        from app.services.email_service import EmailSender

        mock_smtp = MagicMock()
        monkeypatch.setattr(smtplib, "SMTP", lambda *a, **kw: mock_smtp)

        sender = EmailSender(config={"server": "smtp.test.com", "port": 587, "user": "u", "password": "p"})
        result = sender.send_custom_email("dest1@test.com, dest2@test.com", "Assunto", "<b>Html</b>")
        assert result is True

    def test_export_pdf_with_special_characters_and_urls(self, auth_client):
        # 1. Export PDF de empresas com & e caracteres especiais
        special_companies = [
            {
                "id": 1,
                "nome": "COMÉRCIO & INDÚSTRIA LTDA <MATRIZ>",
                "cnpj": "11.222.333/0001-44",
                "origem": "Manual & Planilha",
                "status": True
            },
            {
                "id": 2,
                "nome": "ALIMENTOS & CIA S.A.",
                "cnpj": "22.333.444/0001-55",
                "origem": "GestãoClick",
                "status": False
            }
        ]
        resp_comp = auth_client.post('/api/export_pdf', json={"companies": special_companies})
        assert resp_comp.status_code == 200
        assert resp_comp.mimetype == 'application/pdf'
        assert len(resp_comp.data) > 0

        # 2. Export PDF de menções com & no nome, trecho com tags e links com query params
        special_mentions = [
            {
                "id": "spec1",
                "empresa": "DISTRIBUIDORA DE BEBIDAS & ALIMENTOS LTDA",
                "cnpj": "11.222.333/0001-44",
                "secao": "DOU - Seção 1 & 2",
                "data": "25/08/2026",
                "trecho": "Publicação referente à empresa <b>DISTRIBUIDORA & CIA</b> com CNPJ 11.222.333/0001-44 e valor de R$ 50.000,00 > R$ 30.000,00",
                "link": "https://in.gov.br/dou?art_id=12345&secao=1&termo=bebidas&data=25/08/2026"
            }
        ]
        resp_mentions = auth_client.post('/api/export_mentions_pdf', json={"mentions": special_mentions})
        assert resp_mentions.status_code == 200
        assert resp_mentions.mimetype == 'application/pdf'
        assert len(resp_mentions.data) > 0

    def test_toggle_main_routine_updates_sqlite(self, auth_client, app):
        from app.models import Settings

        # Pausa a rotina principal
        resp = auth_client.post('/api/routines/toggle/Pesquisa_cnpj.yaml', json={"active": False})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data.get('active') is False

        with app.app_context():
            s_rec = Settings.query.filter_by(key='main_dag_settings').first()
            assert s_rec is not None
            val = s_rec.get_value()
            assert val.get('active') is False

        # Reativa a rotina principal
        resp2 = auth_client.post('/api/routines/toggle/Pesquisa_cnpj.yaml', json={"active": True})
        assert resp2.status_code == 200
        data2 = resp2.get_json()
        assert data2.get('active') is True

        with app.app_context():
            s_rec2 = Settings.query.filter_by(key='main_dag_settings').first()
            assert s_rec2.get_value().get('active') is True












